"""Fictional layout regression: split name, summary sidebar, aligned job rows.

No uploaded document or personal contact data is embedded in this fixture.
"""

from __future__ import annotations

import io

CONTACT = ["Avery", "Ramos", "Logistics Operations Specialist", "avery.ramos@example.com",
           "+1 312 555 0199", "Chicago, IL", "linkedin.com/in/avery-ramos-example"]
SUMMARY = [
    "Supply chain professional with experience in inbound logistics,",
    "procurement, inventory planning, and supplier performance across domestic",
    "and international operations. Experienced with freight providers,",
    "inventory records and cross-functional process improvement.",
]
JOBS = [
    ("Senior Supply Chain Analyst", "Cedar Freight - Evanston, IL", "Apr 2024 - Present", 6),
    ("Supply Chain Analyst", "Harbor Retail - Chennai, India", "Jun 2020 - Nov 2022", 5),
]


def sidebar_resume_bytes(fmt: str, *, contact_on_right: bool = False, labelled_summary: bool = False) -> bytes:
    buffer = io.BytesIO()
    education = [
        ["Master's in Industrial Engineering", "Cascadia Institute of Technology", "Jan 2023 - Dec 2024"],
        ["and Operations", "Chicago, IL", ""],
    ]
    certifications = ["Example Logistics Credential", "Example Inventory Credential"]
    if fmt == "docx":
        from docx import Document

        document = Document()
        cells = document.add_table(rows=1, cols=2).rows[0].cells
        cells[int(contact_on_right)].text = "\n".join(CONTACT)
        cells[1 - int(contact_on_right)].text = "\n".join((["Professional Summary"] if labelled_summary else []) + SUMMARY)
        document.add_paragraph("Work Experience")
        for title, employer, dates, count in JOBS:
            cells = document.add_table(rows=1, cols=3).rows[0].cells
            for cell, value in zip(cells, (title, employer, dates)):
                cell.text = value
            for i in range(count):
                document.add_paragraph(f"Coordinated {employer.split(' - ')[0]} shipment workflow {i + 1}.", style="List Bullet")
        document.add_paragraph("Education")
        table = document.add_table(rows=2, cols=3)
        for row, values in zip(table.rows, education):
            for cell, value in zip(row.cells, values):
                cell.text = value
        for line in ["Technical Skills", "Planning: Inventory management, SQL, demand planning", "Certifications", *certifications]:
            document.add_paragraph(line)
        document.save(buffer)
        return buffer.getvalue()

    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=792, height=900)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica"), NameObject("/Encoding"): NameObject("/WinAnsiEncoding")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    commands: list[str] = []

    def draw(x: int, y: int, value: str, size: int = 9) -> None:
        escaped = value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"BT /F1 {size} Tf {x} {y} Td ({escaped}) Tj ET")

    contact_x, summary_x = (520, 40) if contact_on_right else (40, 290)
    for i, line in enumerate(CONTACT):
        draw(contact_x, 850 - i * 22, line, 20 if i < 2 else 9)
    if labelled_summary:
        draw(summary_x, 862, "Professional Summary")
    for i, line in enumerate(SUMMARY):
        draw(summary_x, 848 - i * 13, line)
    y = 670
    draw(40, y, "Work Experience")
    for title, employer, dates, count in JOBS:
        y -= 25
        for x, value in zip((40, 280, 600), (title, employer, dates)):
            draw(x, y, value)
        for i in range(count):
            y -= 18
            draw(40, y, f"- Coordinated {employer.split(' - ')[0]} shipment workflow {i + 1}.")
    y -= 30
    draw(40, y, "Education")
    for row in education:
        y -= 18
        for x, value in zip((40, 280, 600), row):
            # Separate runs emulate a font change inside a degree, not two fields.
            if value.startswith("Master's in"):
                draw(x, y, "Master's")
                draw(x + 35, y, value[len("Master's"):])
            elif value:
                draw(x, y, value)
    for line in ["Technical Skills", "Planning: Inventory management, SQL, demand planning", "Certifications", *certifications]:
        y -= 18
        draw(40, y, line)
    stream = DecodedStreamObject()
    stream.set_data("\n".join(commands).encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(buffer)
    return buffer.getvalue()
