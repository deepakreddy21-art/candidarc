"""Synthetic supply-chain layout; no private résumé or contact data."""

from __future__ import annotations

import io

PAGE_ONE = """Casey Morgan
Chicago, IL | +1 708 555 0101 | casey@example.com
LinkedIn | Portfolio
SUMMARY
Supply chain analyst supporting international vendors.
TECHNICAL SKILLS
ERP & Tools: Netsuite, SAP MM, WMS
Forecasting & Planning: demand planning, safety stock modeling,
inventory optimization
Data & Analysis: Excel (advanced formulas, pivot tables,
Power Query, VLOOKUP), SQL, KPI dashboards,
inventory turnover tracking
Process & Soft Skills: stakeholder communication, continuous
improvement, process automation
EXPERIENCE
Harbor Logistics     Bensenville, Illinois
Senior Supply Chain Analyst     May 2024 - Present
- Led inventory forecasting for international vendors.
- Negotiated contracts for recurring purchases.
Market Labs     Hyderabad, India
Supply Chain Analyst     May 2020 - Dec 2022
- Managed purchase orders and shipment records.
- Coordinated customs documentation for suppliers.
EDUCATION
Master's in Industrial Technology and Operations     Chicago, IL"""

PAGE_TWO = """Lakeside Institute of Technology
CERTIFICATIONS
ASCM: Certified Supply Chain Professional
ISM (Institute for Supply Management): Certified Professional in Supply Management"""

LINKS = {"LinkedIn": "https://www.linkedin.com/in/casey-example", "Portfolio": "https://casey.example.com"}


def linked_resume_bytes(fmt: str = "pdf") -> bytes:
    buffer = io.BytesIO()
    if fmt == "docx":
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        document = Document()
        for page in (PAGE_ONE, PAGE_TWO):
            for line in page.splitlines():
                paragraph = document.add_paragraph()
                if line == "LinkedIn | Portfolio":
                    for label, url in LINKS.items():
                        hyperlink = OxmlElement("w:hyperlink")
                        hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                        run, text = OxmlElement("w:r"), OxmlElement("w:t")
                        text.text = label
                        run.append(text)
                        hyperlink.append(run)
                        paragraph._p.append(hyperlink)
                        paragraph.add_run(" | ")
                else:
                    paragraph.add_run(line)
            if page == PAGE_ONE:
                document.add_page_break()
        document.save(buffer)
        return buffer.getvalue()

    from pypdf import PdfWriter
    from pypdf.annotations import Link
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for page_text in (PAGE_ONE, PAGE_TWO):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica"), NameObject("/Encoding"): NameObject("/WinAnsiEncoding")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        operators: list[str] = []
        for index, line in enumerate(page_text.splitlines()):
            y = 750 - index * 18
            cells = [(40, line)]
            if line == "LinkedIn | Portfolio":
                cells = [(40, "LinkedIn"), (150, "Portfolio")]
                for x, label in cells:
                    writer.add_annotation(len(writer.pages) - 1, Link(rect=(x, y - 2, x + 65, y + 12), url=LINKS[label]))
            for x, value in cells:
                escaped = value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                operators.append(f"BT /F1 10 Tf {x} {y} Td ({escaped}) Tj ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(operators).encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(buffer)
    return buffer.getvalue()
