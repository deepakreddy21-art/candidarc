"""Fictional finance résumé: compound titles, table skills and subsection captions."""

from __future__ import annotations

import io

SKILL_ROWS = [
    ("Financial Planning", "FP&A, Budgeting, Revenue Forecasting"),
    ("Inventory & Costing", "Inventory Accounting, Standard Costing, Reconciliations"),
    ("Reporting & Analytics", "Management Reporting, KPI Design & Monitoring"),
    ("Systems & Tools", "Excel (Pivot Tables, Power Query, XLOOKUP), SQL, Oracle Hyperion, SAP FICO"),
    ("Controls & Audit", "Internal Controls, Audit Documentation, Version Control"),
    ("Collaboration", "Stakeholder Communication, Process Improvement"),
]

CAREER = """PROFESSIONAL EXPERIENCE
Operations / Financial Analyst (Contract) | Cedar Instruments\tApr 2024 – Present | Chicago, IL
Inventory & Month-End Close Ownership
- Reconciled Cedar inventory balances with the general ledger.
- Prepared Cedar inventory schedules for monthly review.
Financial Modeling & Forecasting
- Built Cedar budget forecasts for component purchases.
- Reviewed Cedar supplier pricing scenarios.
Automation, Controls & Executive Reporting
- Automated Cedar variance reports with SQL.
- Documented Cedar controls for recurring close tasks.
- Presented Cedar operating expense trends.
Strategic Financial Analyst Intern | Fairway Markets\tJun 2023 – Dec 2023 | Chicago, IL
Retail Inventory Analysis & Automation
- Analyzed Fairway stock counts and reconciliation records.
- Built Fairway inventory dashboards in Excel.
Variance Analysis & Process Redesign
- Reviewed Fairway budget variances with store managers.
- Documented Fairway report validation steps.
Executive – Finance & Operations | Aster Aviation (Regional Air)\tAug 2018 – Nov 2022 | Hyderabad, India
Inventory Finance & Standard Costing
- Prepared Aster aviation cost schedules.
- Reconciled Aster spare parts ledgers.
Physical Counts, Dashboards & Internal Controls
- Coordinated Aster physical inventory counts.
- Maintained Aster control documentation.
- Prepared Aster operations dashboards.
KEY PROJECTS
Budget Scenario Planner
- Built an Excel scenario model for inventory planning.
- Added validation checks to the model.
Inventory Finance Dashboard
- Built a SQL report for stock reconciliation.
Valuation Research Model
- Compared discounted cash-flow scenarios.
EDUCATION
Master of Science, Finance | Eastlake University, Naperville, IL\tSep 2022 – May 2024
Thesis: Budget planning for seasonal inventory
Bachelor of Commerce | Marina College, Chennai, India\tJun 2015 – Apr 2018
"""


def finance_resume_bytes(fmt: str) -> bytes:
    buffer = io.BytesIO()
    contact = ["Casey Morgan", "casey.finance@example.com | +1 708 555 0101 | Chicago, IL"]
    if fmt == "docx":
        from docx import Document

        document = Document()
        for line in [*contact, "CORE COMPETENCIES"]:
            document.add_paragraph(line)
        table = document.add_table(rows=0, cols=2)
        for category, skills in SKILL_ROWS:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = category, skills
        for line in CAREER.splitlines():
            if line.startswith("- "):
                document.add_paragraph(line[2:], style="List Bullet")
            else:
                document.add_paragraph(line)
        document.save(buffer)
        return buffer.getvalue()

    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    # Real font-encoded PDF with spaced headings, aligned cells and wrapped place.
    lines = [*contact, "C O R E   C O M P E T E N C I E S"]
    lines.extend(f"{category}      {skills}" for category, skills in SKILL_ROWS)
    career = CAREER.replace("Chicago, IL\nInventory &", "Chicago,\nIL\nInventory &", 1)
    lines.extend(" ".join(line) if line in {"PROFESSIONAL EXPERIENCE", "KEY PROJECTS", "EDUCATION"}
                 else line.replace("\t", "     ") for line in career.splitlines())
    writer = PdfWriter()
    for start in range(0, len(lines), 32):
        page = writer.add_blank_page(width=792, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica"), NameObject("/Encoding"): NameObject("/WinAnsiEncoding")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        operators: list[str] = []
        for index, line in enumerate(lines[start:start + 32]):
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            operators.append(f"BT /F1 9 Tf 35 {750 - index * 21} Td ({escaped}) Tj ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(operators).encode("cp1252"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(buffer)
    return buffer.getvalue()
