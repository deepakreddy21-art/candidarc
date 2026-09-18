"""Record-level import regressions: layouts, not a list of known employers/cities."""

from __future__ import annotations

import base64
import io

import pytest

from app.modules.parsing.service import parse_resume_bytes_sync
from app.modules.parsing.structure import structure_resume_text
from tests.fixtures.resume_samples import text_to_docx, text_to_simple_pdf


@pytest.mark.parametrize("heading", [
    "WORK EXPERIENCE", "Professional Experience", "Employment History", "Career History",
    "PROFESSIONAL BACKGROUND", "Relevant Work Experience", "01. Work & Professional Experience",
])
@pytest.mark.parametrize("header", [
    "Harbor Mutual Software Engineer San Antonio, TX, USA",
    "Harbor Mutual | Software Engineer | San Antonio, TX, USA",
    "Software Engineer | Harbor Mutual | San Antonio, TX, USA",
    "Harbor Mutual\nSoftware Engineer\nSan Antonio, TX, USA",
    "Harbor Mutual    San Antonio, TX, USA\nSoftware Engineer",
    "Software Engineer at Harbor Mutual, San Antonio, TX, USA",
])
def test_employment_order_and_section_aliases(heading: str, header: str):
    parsed = structure_resume_text(
        f"Alex Rivera\n{heading}\n{header}\nJan 2024 - Present\n"
        "- Built REST APIs with Python and Java.\n"
        "Maintained these services across multiple regions.\n"
        "EDUCATION\nB.S. Computer Science | Lakeside University | 2020"
    )
    assert len(parsed["employment"]) == 1
    row = parsed["employment"][0]
    assert (row["title"], row["employer"], row["location"]) == (
        "Software Engineer", "Harbor Mutual", "San Antonio, TX, USA",
    )
    assert (row["start_date"], row["end_date"], row["is_current"]) == ("Jan 2024", "Present", True)
    assert len(row["bullets"]) == 1
    assert "Maintained these services" in row["bullets"][0]


@pytest.mark.parametrize("location", ["Zürich, Switzerland", "Hyderabad, India", "São Paulo, Brazil", "Montréal, Québec, Canada", "Remote"])
def test_locations_are_not_a_us_city_allowlist(location: str):
    parsed = structure_resume_text(
        f"Alex Rivera\nWORK HISTORY\nExample Labs | ML Engineer | {location}\n"
        "2021 - 2024\n- Built data pipelines."
    )
    row = parsed["employment"][0]
    assert row["title"] == "ML Engineer"
    assert row["employer"] == "Example Labs"
    assert row["location"] == location


@pytest.mark.parametrize("education", [
    "Lakeside Institute of Technology | Chicago, ILJan 2023 - May 2024\nMaster of Science in Information Technology",
    "Lakeside Institute of Technology    Chicago, IL\nMaster of Science in Information Technology    Jan 2023 - May 2024",
    "Master of Science in Information Technology\nLakeside Institute of Technology\nChicago, IL\nJan 2023 - May 2024",
    "Lakeside Institute of Technology\nJan 2023 to May 2024\nChicago, IL\nMaster of Science in Information Technology",
])
def test_multiline_education_is_one_record(education: str):
    parsed = structure_resume_text(f"Alex Rivera\nACADEMIC QUALIFICATIONS\n{education}\nSKILLS\nPython, SQL")
    assert len(parsed["education"]) == 1
    row = parsed["education"][0]
    assert row["institution"] == "Lakeside Institute of Technology"
    assert row["degree"] == "Master of Science"
    assert row["field"] == "Information Technology"
    assert row["location"] == "Chicago, IL"
    assert (row["start_date"], row["end_date"]) == ("Jan 2023", "May 2024")


@pytest.mark.parametrize("degree,field", [
    ("Ph.D.", "Computer Science"), ("B.Tech", "Computer Engineering"),
    ("Bachelor of Engineering", "Mechanical Engineering"), ("Master’s", "Data Science"),
])
def test_education_degree_before_or_after_acronym_institution(degree: str, field: str):
    parsed = structure_resume_text(
        f"Alex Rivera\nEDUCATION\nMIT\n{degree} in {field}\nCambridge, MA\n2018 - 2022\n"
        "Stanford University\nM.S. in Statistics\n2022 - 2024\nSKILLS\nSQL"
    )
    assert len(parsed["education"]) == 2
    one, two = parsed["education"]
    assert (one["institution"], one["degree"], one["field"]) == ("MIT", degree, field)
    assert one["location"] == "Cambridge, MA"
    assert two["institution"] == "Stanford University"
    assert two["field"] == "Statistics"


@pytest.mark.parametrize("heading", ["Project Experience", "Academic Projects", "Selected Technical Projects", "Research Projects", "Personal Projects"])
def test_projects_keep_multiline_bullets_dates_and_own_section(heading: str):
    parsed = structure_resume_text(
        f"Alex Rivera\nEDUCATION\nB.S. Computer Science | Cascadia University | 2020\n{heading}\n"
        "Atlas Scheduler | Lead Developer | Campus Lab | Jan 2023 - May 2023\n"
        "- Built a scheduling system with Python\n"
        "and PostgreSQL for the campus lab.\n"
        "https://github.com/example/atlas\n"
        "Beacon Monitor | Jun 2023 - Aug 2023\n- Built a React dashboard.\n"
        "SKILLS\nPython, React\nAWARDS\nDean's Award 2020"
    )
    assert len(parsed["education"]) == 1
    assert len(parsed["projects"]) == 2
    first, second = parsed["projects"]
    assert (first["name"], first["role"], first["organization"]) == ("Atlas Scheduler", "Lead Developer", "Campus Lab")
    assert first["start_date"] == "Jan 2023"
    assert first["end_date"] == "May 2023"
    assert len(first["bullets"]) == 1
    assert "and PostgreSQL" in first["bullets"][0]
    assert first["repo_url"] == "https://github.com/example/atlas"
    assert second["name"] == "Beacon Monitor"
    assert parsed["skills"] == ["Python", "React"]


def test_unknown_title_retains_source_and_is_marked_for_review():
    parsed = structure_resume_text(
        "Alex Rivera\nWORK EXPERIENCE\nOrbital Wrangler | Meridian Labs | Remote\n"
        "2020 - 2024\n- Operated the platform.\nSKILLS\nPython, SQL"
    )
    assert len(parsed["employment"]) == 1
    row = parsed["employment"][0]
    assert row["provenance"]["confidence"] != "high"
    assert row["provenance"]["warnings"]
    assert "Orbital Wrangler" in row["provenance"]["source_text"]
    assert parsed["extraction_quality"] != "high"


@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_real_bytes_map_collapsed_job_header_and_multiline_school(fmt: str):
    text = (
        "Alex Rivera\nalex@example.com\nWORK EXPERIENCE\n"
        "USAA Software Engineer San Antonio, TX\nJan 2024 - Present\n- Built Python APIs.\n"
        "EDUCATION\nIllinois Institute of Technology | Chicago, IL\nJan 2023 - May 2024\n"
        "Master of Science in Information Technology\nSKILLS\nPython, SQL"
    )
    raw = text_to_simple_pdf(text) if fmt == "pdf" else text_to_docx(text)
    result = parse_resume_bytes_sync(f"layout.{fmt}", "application/pdf" if fmt == "pdf" else
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                    base64.b64encode(raw).decode())
    assert len(result.employment) == len(result.education) == 1
    assert result.employment[0].title == "Software Engineer"
    assert result.employment[0].employer == "USAA"
    assert result.employment[0].location == "San Antonio, TX"
    assert result.education[0].institution == "Illinois Institute of Technology"
    assert result.education[0].degree == "Master of Science"
    assert result.education[0].field == "Information Technology"
    assert result.education[0].end_date == "May 2024"


def test_pdf_positioned_header_cells_and_right_aligned_dates_stay_in_their_record():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    # Valid xref and independently positioned cells, not flattened text in a PDF.
    items = [
        (40, 760, "Alex Rivera"), (40, 740, "alex@example.com"),
        (40, 715, "WORK EXPERIENCE"),
        (40, 690, "Harbor Mutual"), (205, 690, "Software Engineer"), (430, 690, "San Antonio, TX"),
        (430, 675, "Jan 2024 - Present"), (40, 650, "- Built Python APIs."),
        (40, 620, "EDUCATIONAL BACKGROUND"),
        (40, 600, "Lakeside Institute of Technology"), (430, 600, "Chicago, IL"),
        (40, 580, "Master of Science in Information Technology"), (430, 580, "Jan 2023 - May 2024"),
        (40, 550, "PROJECT EXPERIENCE"), (40, 530, "Atlas Scheduler"),
        (430, 530, "2022 - 2023"), (40, 510, "- Built SQL reports."),
    ]
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data("\n".join(f"BT /F1 9 Tf {x} {y} Td ({text}) Tj ET" for x, y, text in items).encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    result = parse_resume_bytes_sync("positioned.pdf", "application/pdf", base64.b64encode(buffer.getvalue()).decode())
    assert len(result.employment) == len(result.education) == len(result.projects) == 1
    job = result.employment[0]
    assert (job.title, job.employer, job.location, job.start_date) == ("Software Engineer", "Harbor Mutual", "San Antonio, TX", "Jan 2024")
    edu = result.education[0]
    assert (edu.institution, edu.degree, edu.field, edu.location, edu.start_date, edu.end_date) == (
        "Lakeside Institute of Technology", "Master of Science", "Information Technology", "Chicago, IL", "Jan 2023", "May 2024",
    )
    assert result.projects[0].name == "Atlas Scheduler"
    assert result.projects[0].end_date == "2023"


def test_multiple_roles_under_shared_employer_keep_separate_dates():
    result = structure_resume_text(
        "Alex Rivera\nEMPLOYMENT\nExample Labs\nSenior Engineer\n2023 - Present\n- Led platform development.\n"
        "Engineer\n2020 - 2023\n- Built Python APIs."
    )
    assert len(result["employment"]) == 2
    first, second = result["employment"]
    assert first["employer"] == second["employer"] == "Example Labs"
    assert (first["title"], second["title"]) == ("Senior Engineer", "Engineer")
    assert (first["start_date"], second["start_date"]) == ("2023", "2020")


def test_company_and_acronym_school_with_international_location():
    result = structure_resume_text(
        "Alex Rivera\nWORK EXPERIENCE\nJava Developer | Meridian, Hyderabad, India\n2019 - 2022\n- Built Java services.\n"
        "EDUCATION\nBachelor’s | Computer Science | RSTU, Hyderabad, India"
    )
    assert result["employment"][0]["employer"] == "Meridian"
    assert result["employment"][0]["location"] == "Hyderabad, India"
    assert result["education"][0]["institution"] == "RSTU"
    assert result["education"][0]["field"] == "Computer Science"
    assert result["education"][0]["location"] == "Hyderabad, India"


def test_inline_section_headers_do_not_swallow_other_sections():
    result = structure_resume_text(
        "Alex Rivera\nLocation: Zürich, Switzerland\n"
        "Professional Summary: Software engineer focused on reliability.\n"
        "Technical Skills: Python, SQL\n"
        "Project Experience: Atlas Scheduler\n- Built Python APIs.\n"
        "Educational Background: M.S. Statistics | Example University | 2022"
    )
    assert result["contact"]["location"] == "Zürich, Switzerland"
    assert result["professional_summary"] == "Software engineer focused on reliability."
    assert result["skills"] == ["Python", "SQL"]
    assert result["projects"][0]["name"] == "Atlas Scheduler"
    assert result["education"][0]["institution"] == "Example University"


def test_docx_table_text_is_not_discarded():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Alex Rivera")
    doc.add_paragraph("WORK EXPERIENCE")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Harbor Mutual\nSoftware Engineer"
    table.cell(0, 1).text = "San Antonio, TX\nJan 2024 - Present"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "- Built Python APIs."
    doc.add_paragraph("ACADEMIC BACKGROUND")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Lakeside Institute of Technology"
    table.cell(0, 1).text = "Chicago, IL"
    table.cell(1, 0).text = "Master of Science in Information Technology"
    table.cell(1, 1).text = "Jan 2023 - May 2024"
    stream = io.BytesIO()
    doc.save(stream)
    result = parse_resume_bytes_sync("table.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                    base64.b64encode(stream.getvalue()).decode())
    assert len(result.employment) == len(result.education) == 1
    assert result.employment[0].employer == "Harbor Mutual"
    assert result.employment[0].location == "San Antonio, TX"
    assert result.employment[0].bullets == ["Built Python APIs."]
    assert result.education[0].institution == "Lakeside Institute of Technology"
    assert result.education[0].degree == "Master of Science"
    assert result.education[0].end_date == "May 2024"
